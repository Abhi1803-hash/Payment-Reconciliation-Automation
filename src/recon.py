"""
Daily three-way payment reconciliation.

    gateway ledger ─┐
                    ├─► Level 1: payment-level match (exact ID, then normalised ID)
    settlement file ┤
                    ├─► Level 2: UTR batch total vs bank credit (UTR, then date+amount)
    bank statement ─┘
                    └─► Level 3: unexplained bank lines → Claude triage → owner + next action

Output: exception register (Excel), Power BI-ready CSVs, Markdown summary, chart.
Run:    python -m src.recon --as-of 2026-10-15
"""
from __future__ import annotations

import argparse
import re
import sqlite3
import time
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.ai_triage import triage

ROOT = Path(__file__).resolve().parents[1]
RAW, OUT = ROOT / "data" / "raw", ROOT / "output"
DB = ROOT / "data" / "recon.db"
SQL = ROOT / "sql" / "recon.sql"
UTR_RE = re.compile(r"(UTR\d{13,})")

PAYMENT_ACTIONS = {
    "NOT_SETTLED":          ("Settlement Ops", "Raise missing-settlement claim with acquirer"),
    "SHORT_SETTLED":        ("Settlement Ops", "Claim shortfall from acquirer with payment-level proof"),
    "EXCESS_FEE":           ("Finance / Pricing", "Dispute fee vs contracted MDR; claim excess + GST"),
    "OVER_SETTLED":         ("Settlement Ops", "Confirm with acquirer; hold excess in suspense"),
    "DUPLICATE_SETTLEMENT": ("Settlement Ops", "Request reversal of duplicate; block payout of excess"),
    "LATE_SETTLEMENT":      ("Settlement Ops", "Track SLA breach; claim interest/penalty per agreement"),
    "ID_FORMAT_MISMATCH":   ("Data Engineering", "Ask acquirer to fix ID format; no money at risk"),
    "NOT_IN_LEDGER":        ("Data Engineering", "Check for ledger sync gap; trace payment in source system"),
    "PENDING_NOT_DUE":      ("—", "Not due yet; no action"),
}
MONEY_AT_RISK = {"NOT_SETTLED", "SHORT_SETTLED", "EXCESS_FEE", "DUPLICATE_SETTLEMENT", "OVER_SETTLED"}


def load_sql_blocks() -> list[tuple[str, str]]:
    parts = re.split(r"^-- name: (\w+)\s*$", SQL.read_text(encoding="utf-8"), flags=re.M)
    return [(parts[i], parts[i + 1]) for i in range(1, len(parts), 2)]


def run_matching(as_of: str) -> None:
    date.fromisoformat(as_of)                                    # validate before inlining into SQL
    ledger = pd.read_csv(RAW / "gateway_ledger.csv")
    settle = pd.read_csv(RAW / "settlement_report.csv")
    bank = pd.read_csv(RAW / "bank_statement.csv")
    bank["utr_extracted"] = bank.narration.str.extract(UTR_RE, expand=False)
    with sqlite3.connect(DB) as con:
        ledger.to_sql("ledger", con, if_exists="replace", index=False)
        settle.to_sql("settlement_report", con, if_exists="replace", index=False)
        bank.to_sql("bank_statement_parsed", con, if_exists="replace", index=False)
        con.executescript("""
            CREATE INDEX ix_l ON ledger(payment_id);
            CREATE INDEX ix_s ON settlement_report(payment_id);
            CREATE INDEX ix_su ON settlement_report(utr);
            CREATE INDEX ix_b ON bank_statement_parsed(utr_extracted);""")
        for name, block in load_sql_blocks():
            con.executescript(block.replace(":as_of", f"'{as_of}'"))
            if name == "normalise_ids":
                con.execute("CREATE INDEX ix_sn ON settlement_n(pid_norm)")


def ageing_bucket(days: float) -> str:
    if pd.isna(days):
        return "n/a"
    return "0-2 days" if days <= 2 else "3-7 days" if days <= 7 else "8-15 days" if days <= 15 else "15+ days"


def collect(as_of: str) -> dict[str, pd.DataFrame]:
    with sqlite3.connect(DB) as con:
        pay = pd.read_sql("SELECT * FROM payment_exceptions", con)
        batch = pd.read_sql("SELECT * FROM batch_recon", con)
        bank = pd.read_sql("SELECT * FROM unexplained_bank", con)
        totals = pd.read_sql("""
            SELECT (SELECT COUNT(*) FROM ledger) AS ledger_payments,
                   (SELECT ROUND(SUM(expected_net),2) FROM ledger) AS ledger_value,
                   (SELECT COUNT(*) FROM payment_match) AS matched_payments,
                   (SELECT COUNT(*) FROM payment_match WHERE match_type='NORMALISED_ID') AS fuzzy_matched,
                   (SELECT COUNT(DISTINCT utr) FROM settlement_report) AS utr_batches,
                   (SELECT COUNT(*) FROM bank_statement_parsed) AS bank_lines""", con).iloc[0]
        totals = pd.Series({k: (v if k == "ledger_value" else int(v)) for k, v in totals.items()}, dtype=object)

    pay["owner"] = pay.category.map(lambda c: PAYMENT_ACTIONS[c][0])
    pay["next_action"] = pay.category.map(lambda c: PAYMENT_ACTIONS[c][1])
    pay["amount_at_risk"] = pay.apply(
        lambda r: abs(r["diff"]) if r["category"] in MONEY_AT_RISK else 0.0, axis=1).round(2)
    pay["age_days"] = (pd.Timestamp(as_of) - pd.to_datetime(pay.expected_settlement_date)).dt.days
    pay["ageing"] = pay.age_days.map(ageing_bucket)
    pay["priority"] = pd.cut(pay.amount_at_risk, [-1, 0, 500, 5000, float("inf")],
                             labels=["P4 - no money at risk", "P3", "P2", "P1"]).astype(str)
    pay = pay.sort_values(["amount_at_risk", "age_days"], ascending=False, ignore_index=True)

    triaged, ai_source = triage(bank)
    return {"payments": pay, "batches": batch, "bank": triaged, "totals": totals, "ai_source": ai_source}


def score_against_ground_truth(r: dict) -> pd.DataFrame | None:
    """Measure detection accuracy on the injected breaks (synthetic data only)."""
    gt_p, gt_b, gt_k = (RAW / f"_ground_truth_{x}.csv" for x in ("payments", "batches", "bank"))
    if not gt_p.exists():
        return None
    rows = []
    found = r["payments"][r["payments"].category != "PENDING_NOT_DUE"]
    truth = pd.read_csv(gt_p)
    name_map = {"ID_MANGLED": "ID_FORMAT_MISMATCH"}
    truth["expected"] = truth.true_break.replace(name_map)
    for cat, grp in truth.groupby("expected"):
        hit = grp.payment_id.isin(found.loc[found.category == cat, "payment_id"]).sum()
        flagged = (found.category == cat).sum()
        rows.append({"level": "Payment", "break_type": cat, "injected": len(grp), "detected": int(hit),
                     "flagged_total": int(flagged), "recall_pct": round(100 * hit / len(grp), 1),
                     "precision_pct": round(100 * hit / flagged, 1) if flagged else None})
    bt = pd.read_csv(gt_b)
    status_for = {"BATCH_TOTAL_MISMATCH": "BATCH_TOTAL_MISMATCH", "UTR_MISSING_IN_NARRATION": "MATCHED_WITHOUT_UTR"}
    for brk, grp in bt.groupby("true_break"):
        st = status_for[brk]
        hit = grp.utr.isin(r["batches"].loc[r["batches"].status == st, "utr"]).sum()
        flagged = (r["batches"].status == st).sum()
        rows.append({"level": "Batch", "break_type": st, "injected": len(grp), "detected": int(hit),
                     "flagged_total": int(flagged), "recall_pct": round(100 * hit / len(grp), 1),
                     "precision_pct": round(100 * hit / flagged, 1) if flagged else None})
    kt = pd.read_csv(gt_k).merge(r["bank"][["txn_ref", "category"]], on="txn_ref", how="left")
    correct = (kt.true_category == kt.category).sum()
    rows.append({"level": "Bank triage", "break_type": "Category assigned correctly", "injected": len(kt),
                 "detected": int(correct), "flagged_total": int(len(r["bank"])),
                 "recall_pct": round(100 * correct / len(kt), 1),
                 "precision_pct": round(100 * correct / len(r["bank"]), 1) if len(r["bank"]) else None})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ outputs
def summary_table(r: dict) -> pd.DataFrame:
    p = r["payments"][r["payments"].category != "PENDING_NOT_DUE"]
    s = (p.groupby(["category", "owner"], as_index=False)
           .agg(exceptions=("payment_id", "count"), amount_at_risk=("amount_at_risk", "sum")))
    return s.sort_values("amount_at_risk", ascending=False, ignore_index=True)


def make_chart(summary: pd.DataFrame, as_of: str) -> Path:
    plt.rcParams.update({"svg.fonttype": "none", "font.size": 10,
                         "axes.spines.top": False, "axes.spines.right": False})
    s = summary.sort_values("exceptions")
    fig, ax = plt.subplots(figsize=(8.5, 3.8))
    colors = ["#C0392B" if v > 0 else "#95A5A6" for v in s.amount_at_risk]
    bars = ax.barh(s.category, s.exceptions, color=colors)
    for b, amt in zip(bars, s.amount_at_risk):
        ax.text(b.get_width() + 0.5, b.get_y() + b.get_height() / 2,
                f"₹{amt:,.0f} at risk" if amt else "no money at risk", va="center", fontsize=9)
    ax.set_xlabel("Exceptions")
    ax.set_xlim(0, s.exceptions.max() * 1.6)
    ax.set_title(f"Payment-level exceptions by type (as of {as_of})")
    fig.tight_layout()
    path = OUT / "exceptions_by_type.svg"
    fig.savefig(path); plt.close(fig)
    return path


def write_excel(r, summary, accuracy, as_of) -> Path:
    path = OUT / f"recon_exceptions_{as_of}.xlsx"
    t = r["totals"]
    pay = r["payments"]
    open_ex = pay[pay.category != "PENDING_NOT_DUE"]
    headline = pd.DataFrame({"Metric": [
        "As-of date", "Payments in ledger", "Auto-matched payments", "   of which via normalised ID",
        "Auto-match rate", "Open payment exceptions", "Amount at risk (₹)", "UTR batches",
        "Batches matched to bank", "Unexplained bank lines triaged", "Triage engine"],
        "Value": [as_of, f"{t.ledger_payments:,}", f"{t.matched_payments:,}", f"{t.fuzzy_matched:,}",
                  f"{100 * (t.ledger_payments - len(open_ex)) / t.ledger_payments:.2f}%",
                  f"{len(open_ex):,}", f"{open_ex.amount_at_risk.sum():,.2f}", f"{t.utr_batches:,}",
                  f"{(r['batches'].status != 'BATCH_NOT_IN_BANK').sum():,}", f"{len(r['bank']):,}",
                  r["ai_source"]]})
    batch_ex = r["batches"][r["batches"].status != "MATCHED"]
    sheets = {"Summary": headline, "By Category": summary, "Payment Exceptions": open_ex,
              "Batch Exceptions": batch_ex, "Bank Triage": r["bank"]}
    if accuracy is not None:
        sheets["Detection Accuracy"] = accuracy
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    fill, font = PatternFill("solid", fgColor="1F3864"), Font(color="FFFFFF", bold=True)
    prio = {"P1": "FFC7CE", "P2": "FFEB9C", "P3": "FFF2CC"}
    for ws in wb.worksheets:
        for c in ws[1]:
            c.fill, c.font = fill, font
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(ws[1], 1):
            col = get_column_letter(i)
            w = max(len(str(x.value or "")) for x in ws[col][:60]) + 2
            ws.column_dimensions[col].width = min(max(w, 10), 55)
    ws = wb["Payment Exceptions"]
    pc = [c.value for c in ws[1]].index("priority") + 1
    for row in range(2, ws.max_row + 1):
        v = str(ws.cell(row, pc).value)[:2]
        if v in prio:
            for c in ws[row]:
                c.fill = PatternFill("solid", fgColor=prio[v])
    wb.save(path)
    return path


def write_markdown(r, summary, accuracy, chart, as_of) -> Path:
    t, pay = r["totals"], r["payments"]
    open_ex = pay[pay.category != "PENDING_NOT_DUE"]
    bank_view = r["bank"][["value_date", "narration", "credit", "debit", "category", "owner", "needs_review"]]
    lines = [
        f"# Reconciliation summary — as of {as_of}", "",
        f"- **{t.ledger_payments:,}** ledger payments, **{t.matched_payments:,}** auto-matched "
        f"({t.fuzzy_matched} only after ID normalisation)",
        f"- **Auto-match rate: {100 * (t.ledger_payments - len(open_ex)) / t.ledger_payments:.2f}%**",
        f"- **{len(open_ex)}** open payment exceptions, **₹{open_ex.amount_at_risk.sum():,.2f}** at risk",
        f"- **{t.utr_batches:,}** UTR batches checked against **{t.bank_lines:,}** bank lines; "
        f"{(r['batches'].status == 'BATCH_TOTAL_MISMATCH').sum()} total mismatches, "
        f"{(r['batches'].status == 'MATCHED_WITHOUT_UTR').sum()} matched without UTR",
        f"- **{len(r['bank'])}** unexplained bank lines triaged by: {r['ai_source']}", "",
        "## Exceptions by type", "", summary.to_markdown(index=False, floatfmt=",.2f"), "",
        f"![Exceptions by type]({chart.name})", "",
        "## Bank lines needing triage", "", bank_view.to_markdown(index=False), "",
    ]
    if accuracy is not None:
        lines += ["## Detection accuracy vs injected breaks", "", accuracy.to_markdown(index=False), ""]
    path = OUT / f"recon_summary_{as_of}.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main(as_of: str) -> None:
    t0 = time.time()
    OUT.mkdir(parents=True, exist_ok=True)
    run_matching(as_of)
    r = collect(as_of)
    summary = summary_table(r)
    accuracy = score_against_ground_truth(r)
    chart = make_chart(summary, as_of)
    xlsx = write_excel(r, summary, accuracy, as_of)
    md = write_markdown(r, summary, accuracy, chart, as_of)
    # Power BI / Tableau-ready flat files
    r["payments"].to_csv(OUT / "bi_payment_exceptions.csv", index=False)
    r["batches"].drop(columns=["narration"]).to_csv(OUT / "bi_batch_recon.csv", index=False)
    r["bank"].to_csv(OUT / "bi_bank_triage.csv", index=False)
    print(f"[done] {xlsx.name}, {md.name}, BI CSVs in {time.time() - t0:.1f}s")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--as-of", default="2026-10-15", help="Reconciliation date YYYY-MM-DD")
    main(ap.parse_args().as_of)
